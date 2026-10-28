"""Builds the draft NU eIRB protocol (retrospective chart review) as .docx.
Edit the SECTIONS text here, then rerun:  python3 build_irb_docx.py
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

VERSION, DATE = "0.1", "2026-10-28"
NU_PURPLE = RGBColor(0x4E, 0x2A, 0x84)
HERE = Path(__file__).parent
BACKGROUND = (HERE / "background_for_irb.md").read_text().strip()

TITLE = ("Contralateral Cervical Nodal Metastasis in HPV-Associated Oropharyngeal Squamous Cell "
         "Carcinoma Treated with Primary Transoral Surgery: A Retrospective Cohort Study")
SHORT = "CONTRA-OPC"

SECTIONS = [
("1. Study Information", [
 f"Full title: {TITLE}",
 f"Short title: {SHORT}",
 "Principal Investigator: Sandeep Samant, MD, Department of Otolaryngology–Head and Neck Surgery, "
 "Feinberg School of Medicine. [Confirm CITI status]",
 "Co-Investigators: [Radiation Oncology], [Neuroradiology], [Head & Neck Pathology], [Biostatistics], "
 "[Resident/Fellow abstractors]. All must have current CITI Human Subjects and HIPAA training.",
 "Funding: Unfunded (departmental). FDA-regulated: No.",
 "Requested review: Expedited, 45 CFR 46.110 Category 5 (research involving materials collected solely "
 "for non-research purposes). Requested: Waiver of informed consent (45 CFR 46.116(f)) and Waiver of "
 "HIPAA Authorization (45 CFR 164.512(i)).",
 "Cancer-related study: route to the Lurie Cancer Center Scientific Review Committee if required by "
 "current policy. [PI to confirm]",
]),
("2. Objectives", [
 "Primary: To determine the rate of occult pathologic contralateral cervical nodal metastasis in patients "
 "with HPV-associated (p16-positive) oropharyngeal squamous cell carcinoma (OPSCC) who underwent bilateral "
 "neck dissection with a radiologically negative contralateral neck, stratified by the indication for "
 "contralateral dissection (base-of-tongue primary, base-of-tongue extension, soft-palate extension, "
 "equivocal contralateral node).",
 "Secondary: (a) To determine whether ipsilateral nodal burden (number of positive nodes, extranodal "
 "extension, level of involvement) predicts contralateral metastasis. (b) To determine the cumulative "
 "incidence of contralateral regional failure after ipsilateral neck dissection, stratified by whether the "
 "contralateral neck received adjuvant radiation. (c) To measure inter-rater reliability of assessing "
 "base-of-tongue and soft-palate involvement and distance from midline. (d) To describe the levels and "
 "size of occult contralateral deposits, the morbidity of contralateral dissection, and outcomes of "
 "salvage after contralateral failure.",
]),
("3. Background and Significance", [BACKGROUND]),
("4. Study Design", [
 "Single-institution retrospective cohort study (medical record review) of patients treated at "
 "Northwestern Memorial Hospital / Northwestern Medicine. No contact with subjects. No change to care. "
 "The review covers records existing from 2015-01-01 through the date of data lock. We request that "
 "follow-up data recorded in the medical record after IRB approval and up to [date, e.g., 2028-12-31] "
 "be included, so that contralateral failure outcomes can be captured. [Confirm with IRB whether this "
 "requires designation as a retrospective/prospective chart review.]",
]),
("5. Inclusion and Exclusion Criteria", [
 "Inclusion: (1) Age ≥18 years at diagnosis. (2) Biopsy-proven squamous cell carcinoma of the oropharynx "
 "(tonsil, base of tongue, soft palate, pharyngeal wall), or unknown primary localized to the oropharynx "
 "at surgery. (3) Primary surgical treatment (transoral or open) with neck dissection at Northwestern "
 "Medicine between 2015-01-01 and [IRB approval date].",
 "p16-negative cases are included as a separate comparison group and are not pooled with p16-positive cases.",
 "Exclusion: (1) Prior head and neck cancer or prior neck radiation or dissection. (2) Distant metastasis "
 "at diagnosis. (3) Treatment of the primary with intent other than cure. (4) Primary surgery done elsewhere "
 "with only adjuvant care at Northwestern, if operative or pathology records are unavailable.",
 "Vulnerable populations: None targeted. Records of prisoners, pregnant women, or decisionally impaired "
 "adults may be incidentally included. No contact occurs, so no additional risk arises.",
]),
("6. Number of Subjects", [
 "Approximately 250–300 records (estimated ≥250 surgically treated OPSCC since 2015, of whom ~75–80 "
 "underwent bilateral neck dissection). Sample size is set by the available cohort. Precision estimates: "
 "with 78 contralateral cN0 dissections and an observed occult rate of 10%, the exact 95% CI is ~4.5–19%.",
]),
("7. Procedures Involved / Data Collection", [
 "Cohort identification: A Northwestern Medicine Enterprise Data Warehouse (NMEDW) query by ICD-10 "
 "(C01, C09.x, C10.x, C14.2, C77.0 with oropharyngeal primary) and CPT codes for transoral oropharyngeal "
 "resection and neck dissection (e.g., 42842–42845, 42890, 42894, 41120–41155, 38720, 38724), "
 "cross-checked against surgeon case logs.",
 "Data abstraction: Trained study team members will abstract data from Epic, pathology reports, imaging, "
 "and radiation treatment planning systems directly into a Northwestern REDCap database. Variables include "
 "demographics (age, sex, smoking history), tumor characteristics (subsite, laterality, base-of-tongue/palate "
 "extension, p16/HPV status, T stage), clinical and pathologic nodal findings for each side of the neck, "
 "adjuvant treatment fields and doses, complications, and recurrence and survival outcomes. See the "
 "attached data dictionary.",
 "Blinded imaging re-review: A head and neck radiologist, blinded to pathology and outcome, will re-read "
 "pre-treatment imaging for contralateral nodal status and the primary tumor's distance from midline.",
 "Inter-rater reliability: Two reviewers will independently classify base-of-tongue and palate involvement "
 "from imaging and operative/clinical documentation.",
 "No biospecimens will be collected. [If pathology slide re-review for micrometastases is planned, add it "
 "here; it uses existing archived slides only.]",
]),
("8. Risks to Subjects", [
 "This study is minimal risk. The only foreseeable risk is a breach of confidentiality. We minimize it by "
 "storing data only in NU REDCap, limiting access to IRB-approved study personnel, coding records with a "
 "study ID, removing direct identifiers from the analytic dataset, and reporting only aggregate results.",
]),
("9. Potential Benefits", [
 "No direct benefit to subjects. The results may help define which patients can safely avoid treatment "
 "of the opposite side of the neck, which could reduce treatment toxicity for future patients.",
]),
("10. Data Management and Confidentiality", [
 "Storage: NU REDCap (HIPAA-compliant, managed by NUCATS/NU IT). The master link between study ID and MRN "
 "will be kept in identifier-flagged REDCap fields, or in a separate access-restricted file on Northwestern "
 "secure storage. It will not be stored on personal devices, personal network storage, or cloud services "
 "not approved by NU.",
 "Analytic dataset: Exported without direct identifiers. Dates converted to intervals (days from surgery). "
 "Analyzed on NU-managed computers or via NU VPN.",
 "Access: PI and IRB-approved study team only. Audit trail via REDCap.",
 "Sharing: Only aggregate results will be published. Any future multi-institutional pooling will use a "
 "limited or de-identified dataset under a Data Use Agreement and an IRB amendment.",
 "Retention: Data will be kept for at least 7 years after study closure per NU policy, then destroyed. The "
 "identifier link will be destroyed at the earliest point consistent with research needs.",
]),
("11. Waiver of Informed Consent — Justification (45 CFR 46.116(f))", [
 "(1) Minimal risk: Record review only. No interaction or intervention.",
 "(2) Impracticable without the waiver: The cohort spans 2015–present and has ~250–300 patients. A "
 "meaningful proportion have died, moved, or been lost to follow-up. Consenting only reachable patients "
 "would bias the cohort toward survivors without recurrence, which is the outcome under study.",
 "(3) Identifiable information is necessary: Records must be linked across surgery, pathology, imaging, "
 "radiation planning, and follow-up encounters, which requires identifiers during abstraction.",
 "(4) Rights and welfare not adversely affected: No clinical decisions are affected, and confidentiality "
 "protections are in place.",
 "(5) Additional information to subjects: Not applicable. Results are aggregate and have no individual "
 "clinical significance.",
]),
("12. Waiver of HIPAA Authorization — Justification (45 CFR 164.512(i)(2)(ii))", [
 "PHI to be used: name and MRN (linkage only); dates of diagnosis, surgery, radiation, recurrence, and "
 "last follow-up/death (converted to intervals); age; pathology, imaging, operative, and radiation "
 "treatment records related to the oropharyngeal cancer.",
 "(a) Minimal risk to privacy: An adequate plan protects identifiers (Section 10). Identifiers will be "
 "destroyed at the earliest opportunity consistent with the research, and PHI will not be reused or "
 "disclosed except as required by law, for oversight, or for IRB-approved research.",
 "(b) Impracticable without the waiver: See Section 11(2).",
 "(c) Impracticable without access to PHI: Linking records across encounters requires identifiers.",
 "Minimum necessary: Only the data elements listed in the data dictionary will be collected.",
]),
("13. Study Timeline", [
 "IRB approval → REDCap build and NMEDW query (weeks 0–6) → abstraction and blinded re-read (weeks 6–16) "
 "→ analysis (weeks 16–18) → dissemination (weeks 18–24). The study remains open for follow-up updates "
 "through [date].",
]),
("14. Attachments to Upload in eIRB", [
 "☐ This protocol   ☐ REDCap data dictionary (03_data_collection/redcap_data_dictionary.csv)   "
 "☐ NMEDW data request form   ☐ Lurie SRC approval (if required)   ☐ CITI certificates for all personnel   "
 "☐ COI disclosures",
]),
]


def main():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    p = doc.add_paragraph(); r = p.add_run("Northwestern University · Feinberg School of Medicine")
    r.bold = True; r.font.color.rgb = NU_PURPLE; r.font.size = Pt(12)
    p = doc.add_paragraph(); r = p.add_run("eIRB+ Protocol — Retrospective Medical Record Review (DRAFT)")
    r.font.color.rgb = NU_PURPLE
    p = doc.add_paragraph(); r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(13)
    p = doc.add_paragraph(f"IRB #: STU________   Version {VERSION}   {DATE}")
    note = doc.add_paragraph()
    nr = note.add_run("DRAFT for PI review. Transfer each section into the current NU HRP-503 / eIRB+ "
                      "smart-form fields. Section headings follow the standard chart-review structure; "
                      "check them against the template downloaded from irb.northwestern.edu.")
    nr.italic = True; nr.font.size = Pt(9)
    for head, paras in SECTIONS:
        h = doc.add_heading(head, level=2)
        for run in h.runs: run.font.color.rgb = NU_PURPLE
        for t in paras:
            for chunk in t.split("\n\n"):
                doc.add_paragraph(chunk.strip())
    for s in doc.sections:
        fp = s.footer.paragraphs[0]
        fp.text = f"{SHORT} IRB Protocol · Version {VERSION} · {DATE} · CONFIDENTIAL DRAFT"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    out = HERE / f"{SHORT}_IRB_Protocol_v{VERSION}_{DATE}.docx"
    doc.save(out); print("wrote", out)


if __name__ == "__main__":
    main()
